from pathlib import Path
import urllib.request
import warnings
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from scipy.cluster.hierarchy import linkage, dendrogram
from sklearn.preprocessing import StandardScaler
from sklearn.cluster import AgglomerativeClustering
from sklearn.metrics import silhouette_score
from docx import Document
from docx.shared import Inches, Pt

warnings.filterwarnings("ignore", category=FutureWarning)
ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "data"
OUT = ROOT / "outputs"
FIG = OUT / "figures"
DATA_DIR.mkdir(exist_ok=True)
FIG.mkdir(parents=True, exist_ok=True)
CSV = DATA_DIR / "Mall_Customers.csv"
DATA_URL = "https://raw.githubusercontent.com/kennedykwangari/Mall-Customer-Segmentation-Data/master/Mall_Customers.csv"
sns.set_theme(style="whitegrid", context="notebook")

def ensure_dataset():
    if CSV.exists():
        return
    print("Dataset not found locally. Attempting download...")
    try:
        urllib.request.urlretrieve(DATA_URL, CSV)
    except Exception as exc:
        raise SystemExit(
            "Could not download the dataset. Please download Mall_Customers.csv from:\n"
            + DATA_URL + "\nand place it in the data folder.\nReason: " + str(exc)
        )

def savefig(name):
    plt.tight_layout()
    plt.savefig(FIG / name, dpi=180, bbox_inches="tight")
    plt.close()

def main():
    ensure_dataset()
    df = pd.read_csv(CSV)
    required = ["CustomerID", "Gender", "Age", "Annual Income (k$)", "Spending Score (1-100)"]
    missing_cols = [c for c in required if c not in df.columns]
    if missing_cols:
        raise ValueError(f"Dataset is missing expected columns: {missing_cols}")

    # 1. Initial data inspection
    info_text = []
    info_text.append(f"Rows: {df.shape[0]}, columns: {df.shape[1]}")
    info_text.append("Columns: " + ", ".join(df.columns.astype(str)))
    info_text.append("Data types:\n" + df.dtypes.to_string())
    info_text.append("First five rows:\n" + df.head().to_string(index=False))
    info_text.append("Last five rows:\n" + df.tail().to_string(index=False))
    info_text.append("Missing values:\n" + df.isna().sum().to_string())
    info_text.append(f"Duplicate rows: {int(df.duplicated().sum())}")
    (OUT / "dataset_inspection.txt").write_text("\n\n".join(info_text), encoding="utf-8")

    # 2. Preprocessing: clean text, check duplicates, handle missing values.
    clean = df.copy()
    clean["Gender"] = clean["Gender"].astype("string").str.strip().str.title()
    clean = clean.drop_duplicates()
    numeric_cols = ["Age", "Annual Income (k$)", "Spending Score (1-100)"]
    categorical_cols = ["Gender"]
    for col in numeric_cols:
        clean[col] = pd.to_numeric(clean[col], errors="coerce")
        clean[col] = clean[col].fillna(clean[col].median())
    clean["Gender"] = clean["Gender"].fillna(clean["Gender"].mode().iloc[0])
    valid_gender = {"Male", "Female"}
    clean.loc[~clean["Gender"].isin(valid_gender), "Gender"] = clean["Gender"].mode().iloc[0]
    clean["Gender_Code"] = clean["Gender"].map({"Female": 0, "Male": 1}).astype(int)

    features = ["Age", "Annual Income (k$)", "Spending Score (1-100)"]
    # IQR outlier flags are retained for discussion; records are not removed.
    outlier_rows = []
    for col in numeric_cols:
        q1, q3 = clean[col].quantile([0.25, 0.75])
        iqr = q3 - q1
        lo, hi = q1 - 1.5 * iqr, q3 + 1.5 * iqr
        mask = (clean[col] < lo) | (clean[col] > hi)
        outlier_rows.append({
            "Feature": col, "Q1": q1, "Q3": q3, "IQR": iqr,
            "Lower_Fence": lo, "Upper_Fence": hi, "Outlier_Count": int(mask.sum())
        })
    outlier_table = pd.DataFrame(outlier_rows)
    outlier_table.to_csv(OUT / "iqr_outlier_summary.csv", index=False)

    # 3. Statistical analysis
    stats = clean[numeric_cols].agg(["count", "mean", "median", "min", "max", "var", "std"]).T
    stats["range"] = clean[numeric_cols].max() - clean[numeric_cols].min()
    stats["MAD_mean"] = clean[numeric_cols].apply(lambda s: np.mean(np.abs(s - s.mean())))
    stats["IQR"] = clean[numeric_cols].quantile(.75) - clean[numeric_cols].quantile(.25)
    stats.to_csv(OUT / "descriptive_statistics.csv")
    clean[numeric_cols + ["Gender_Code"]].corr(method="pearson").to_csv(OUT / "pearson_correlation.csv")

    # 4. EDA charts
    plt.figure(figsize=(8, 5))
    sns.histplot(data=clean, x="Age", bins=15, kde=True)
    plt.title("Age Distribution")
    savefig("01_age_distribution.png")

    plt.figure(figsize=(8, 5))
    sns.histplot(data=clean, x="Annual Income (k$)", bins=15, kde=True)
    plt.title("Annual Income Distribution")
    savefig("02_income_distribution.png")

    plt.figure(figsize=(8, 5))
    sns.histplot(data=clean, x="Spending Score (1-100)", bins=15, kde=True)
    plt.title("Spending Score Distribution")
    savefig("03_spending_score_distribution.png")

    plt.figure(figsize=(7, 5))
    sns.countplot(data=clean, x="Gender", order=clean["Gender"].value_counts().index)
    plt.title("Customer Count by Gender")
    savefig("04_gender_count.png")

    plt.figure(figsize=(8, 5))
    sns.scatterplot(data=clean, x="Annual Income (k$)", y="Spending Score (1-100)", hue="Gender", s=65)
    plt.title("Annual Income vs Spending Score")
    savefig("05_income_vs_spending.png")

    plt.figure(figsize=(8, 5))
    sns.boxplot(data=clean[numeric_cols])
    plt.title("Boxplots of Numeric Features")
    plt.xticks(rotation=15)
    savefig("06_numeric_boxplots.png")

    plt.figure(figsize=(7, 6))
    sns.heatmap(clean[numeric_cols + ["Gender_Code"]].corr(), annot=True, cmap="coolwarm", center=0, fmt=".2f")
    plt.title("Pearson Correlation Heatmap")
    savefig("07_correlation_heatmap.png")

    # 5. Scale numeric features, then fit hierarchical clustering.
    scaler = StandardScaler()
    X = scaler.fit_transform(clean[features])
    Z = linkage(X, method="ward", metric="euclidean")
    plt.figure(figsize=(14, 6))
    dendrogram(Z, truncate_mode="lastp", p=30, leaf_rotation=90, leaf_font_size=8, show_contracted=True)
    plt.title("Hierarchical Agglomerative Clustering Dendrogram (truncated view)")
    plt.xlabel("Merged observations / clusters")
    plt.ylabel("Ward distance")
    savefig("08_dendrogram.png")

    # Compare candidate cluster counts. Pick the best silhouette score among k=2..8.
    score_rows = []
    labels_by_k = {}
    for k in range(2, 9):
        model = AgglomerativeClustering(n_clusters=k, linkage="ward")
        labels = model.fit_predict(X)
        score = silhouette_score(X, labels)
        score_rows.append({"n_clusters": k, "silhouette_score": score})
        labels_by_k[k] = labels
    scores = pd.DataFrame(score_rows)
    scores.to_csv(OUT / "cluster_count_comparison.csv", index=False)
    best_k = int(scores.loc[scores["silhouette_score"].idxmax(), "n_clusters"])
    final_labels = labels_by_k[best_k]
    clean["Cluster"] = final_labels + 1
    clean.to_csv(OUT / "customer_segments.csv", index=False)

    plt.figure(figsize=(8, 5))
    sns.lineplot(data=scores, x="n_clusters", y="silhouette_score", marker="o")
    plt.xticks(range(2, 9))
    plt.title("Silhouette Score by Number of Clusters")
    plt.xlabel("Number of clusters")
    plt.ylabel("Silhouette score")
    savefig("09_silhouette_by_k.png")

    plt.figure(figsize=(9, 6))
    sns.scatterplot(data=clean, x="Annual Income (k$)", y="Spending Score (1-100)", hue="Cluster", palette="tab10", s=75)
    plt.title(f"Customer Segments: Income vs Spending (k={best_k})")
    savefig("10_cluster_scatter.png")

    plt.figure(figsize=(9, 5))
    cluster_sizes = clean["Cluster"].value_counts().sort_index()
    sns.barplot(x=cluster_sizes.index.astype(str), y=cluster_sizes.values)
    plt.title("Number of Customers in Each Cluster")
    plt.xlabel("Cluster")
    plt.ylabel("Customers")
    savefig("11_cluster_sizes.png")

    # 6. Cluster profiles and report
    profile = clean.groupby("Cluster")[features].mean().round(2)
    profile["Customers"] = clean.groupby("Cluster").size()
    profile.to_csv(OUT / "cluster_profiles.csv")

    gender_counts = clean["Gender"].value_counts()
    corr = clean[numeric_cols + ["Gender_Code"]].corr(method="pearson")
    income_spend_corr = float(corr.loc["Annual Income (k$)", "Spending Score (1-100)"])
    age_mean = float(clean["Age"].mean())
    income_mean = float(clean["Annual Income (k$)"].mean())
    spending_mean = float(clean["Spending Score (1-100)"].mean())
    silhouette_best = float(scores.loc[scores["n_clusters"] == best_k, "silhouette_score"].iloc[0])
    duplicates_before = int(df.duplicated().sum())
    missing_before = int(df.isna().sum().sum())

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    doc.add_heading("MACHINE LEARNING MINI-PROJECT", 0)
    p = doc.add_paragraph()
    p.alignment = 1
    p.add_run("Group 12 | Hierarchical Agglomerative Clustering\n").bold = True
    p.add_run("Mall Customer Segmentation")
    doc.add_paragraph("Submitted as a group mini-project for the Machine Learning assignment. Add the group members' names, roll numbers, course, college, and submission date on this page before submission.")
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        f"This project applies hierarchical agglomerative clustering to {len(clean)} mall-customer records to identify groups with similar age, annual income, and spending-score profiles. "
        f"After data inspection, cleaning, statistical analysis, visual exploration, and standardization, Ward-linkage clustering was compared for 2–8 clusters using silhouette score. "
        f"The best score in this tested range was {silhouette_best:.3f} for {best_k} clusters. The resulting groups provide an exploratory basis for customer-focused marketing; they are not ground-truth customer categories."
    )
    doc.add_heading("1. Introduction and Problem Statement", level=1)
    doc.add_paragraph(
        "Retailers serve customers with different demographic characteristics and spending patterns. Treating every customer identically can make campaigns inefficient. "
        "The objective is to investigate whether hierarchical agglomerative clustering can reveal interpretable customer segments from age, annual income, and spending score."
    )
    doc.add_heading("2. Dataset Description", level=1)
    doc.add_paragraph(
        f"The CSV contains {df.shape[0]} rows and {df.shape[1]} columns. Fields are CustomerID, Gender, Age, Annual Income (k$), and Spending Score (1-100). "
        "CustomerID is an identifier and is excluded from distance calculations. Gender is summarized during EDA but the core clustering uses the three numeric customer attributes."
    )
    doc.add_paragraph("Dataset source: https://github.com/kennedykwangari/Mall-Customer-Segmentation-Data/blob/master/Mall_Customers.csv")
    doc.add_heading("3. Methodology", level=1)
    for title, body in [
        ("Data loading", "Loaded the CSV with pandas; inspected shape, columns, types, head/tail, missing values, and duplicate rows."),
        ("Preprocessing", "Trimmed and standardized Gender labels, removed exact duplicate rows, coerced numeric fields, and imputed any numeric missing values with the median and categorical missing values with the mode. Inappropriate gender labels are replaced with the mode."),
        ("Feature scaling", "StandardScaler transforms age, annual income, and spending score to comparable scales. This is important because Euclidean distances can otherwise be dominated by features with larger numerical ranges."),
        ("Outlier detection", "Used the 1.5×IQR fences for each numeric feature. Potential outliers are flagged and reported, not automatically deleted, because unusual customer profiles may be meaningful."),
        ("Hierarchical clustering", "Ward linkage starts with each observation as its own cluster and repeatedly merges the pair of clusters that produces the smallest increase in within-cluster variance. The dendrogram visualizes the merge structure."),
        ("Cluster selection", f"Tested k=2 through k=8 and selected k={best_k} based on the highest silhouette score within that range. This is a data-driven heuristic, not proof that the selected number is uniquely correct."),
        ("Evaluation", "Silhouette score measures how similar each point is to its own cluster compared with other clusters. Clustering is unsupervised, so accuracy, precision, recall, F1-score, and ROC-AUC are not applicable without true labels.")
    ]:
        doc.add_heading(title, level=2)
        doc.add_paragraph(body)
    doc.add_heading("4. Exploratory Data Analysis", level=1)
    doc.add_paragraph(
        f"The average age is {age_mean:.2f} years, average annual income is {income_mean:.2f} thousand dollars, and average spending score is {spending_mean:.2f}. "
        f"Gender counts are {gender_counts.to_dict()}. The Pearson correlation between annual income and spending score is {income_spend_corr:.3f}; this describes linear association only and should not be interpreted as causation."
    )
    for image_name, caption in [
        ("01_age_distribution.png", "Figure 1. Age distribution."),
        ("02_income_distribution.png", "Figure 2. Annual income distribution."),
        ("03_spending_score_distribution.png", "Figure 3. Spending-score distribution."),
        ("04_gender_count.png", "Figure 4. Customer count by gender."),
        ("05_income_vs_spending.png", "Figure 5. Annual income compared with spending score."),
        ("06_numeric_boxplots.png", "Figure 6. Numeric feature boxplots."),
        ("07_correlation_heatmap.png", "Figure 7. Pearson correlation heatmap.")
    ]:
        doc.add_paragraph(caption)
        doc.add_picture(str(FIG / image_name), width=Inches(5.7))
    doc.add_heading("5. Statistical Analysis and Outliers", level=1)
    doc.add_paragraph(
        "The generated descriptive_statistics.csv reports count, mean, median, minimum, maximum, variance, standard deviation, range, mean absolute deviation from the mean, and IQR for the numerical variables. "
        "The IQR outlier table includes quartiles, fences, and counts outside the fences. A flagged point is not automatically an error; domain context is needed before removal."
    )
    doc.add_heading("6. Model Building and Evaluation", level=1)
    doc.add_paragraph(f"The selected model uses Ward linkage with {best_k} clusters. The best silhouette score among k=2–8 is {silhouette_best:.3f}. A higher silhouette score generally indicates better separation and cohesion, but it should be interpreted alongside the dendrogram and cluster profiles.")
    doc.add_paragraph("Figure 8. Truncated hierarchical dendrogram.")
    doc.add_picture(str(FIG / "08_dendrogram.png"), width=Inches(6.0))
    doc.add_paragraph("Figure 9. Silhouette score for candidate cluster counts.")
    doc.add_picture(str(FIG / "09_silhouette_by_k.png"), width=Inches(5.7))
    doc.add_paragraph("Figure 10. Customer segments by annual income and spending score.")
    doc.add_picture(str(FIG / "10_cluster_scatter.png"), width=Inches(5.7))
    doc.add_paragraph("Figure 11. Number of customers in each cluster.")
    doc.add_picture(str(FIG / "11_cluster_sizes.png"), width=Inches(5.7))
    doc.add_heading("7. Cluster Profiles", level=1)
    table = doc.add_table(rows=1, cols=len(profile.columns) + 1)
    table.style = "Light Shading Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Cluster"
    for i, col in enumerate(profile.columns, start=1):
        hdr[i].text = str(col)
    for idx, row in profile.iterrows():
        cells = table.add_row().cells
        cells[0].text = str(idx)
        for i, val in enumerate(row.values, start=1):
            cells[i].text = str(val)
    doc.add_paragraph(
        "Interpret each cluster using its average age, annual income, spending score, and size. Cluster numbers are arbitrary labels; they do not represent ranking or quality. "
        "Marketing actions should be based on the profile and validated with additional customer and business data."
    )
    doc.add_heading("8. Findings and Interpretation", level=1)
    doc.add_paragraph(f"1. The analysis partitions the customers into {best_k} groups under the selected criterion.")
    doc.add_paragraph(f"2. The mean age, income, and spending score are {age_mean:.2f}, {income_mean:.2f} k$, and {spending_mean:.2f}, respectively.")
    doc.add_paragraph(f"3. Income–spending Pearson correlation is {income_spend_corr:.3f}; the scatterplot should be inspected for non-linear or subgroup patterns.")
    doc.add_paragraph("4. Cluster profiles can support exploratory campaign design, loyalty-program experiments, and differentiated customer communication.")
    doc.add_heading("9. Limitations", level=1)
    doc.add_paragraph(
        "The dataset is small and contains only a few customer attributes. Spending score is an assigned score rather than transaction-level revenue. "
        "The clustering result depends on feature selection, scaling, linkage, distance, and the selected number of clusters. There are no true segment labels, so external validation is unavailable. "
        "Clusters should be validated with larger, more recent, and richer data before operational use."
    )
    doc.add_heading("10. Future Improvements", level=1)
    doc.add_paragraph(
        "Compare Ward, average, and complete linkage; test alternative feature sets; examine stability under resampling; compare with K-Means and DBSCAN; and validate segments using transaction history, campaign response, and business outcomes."
    )
    doc.add_heading("11. Conclusion", level=1)
    doc.add_paragraph(
        f"This mini-project implements an end-to-end unsupervised learning workflow and produces {best_k} exploratory customer segments using hierarchical agglomerative clustering. "
        "The dendrogram, silhouette comparison, visualizations, and cluster profiles make the analysis interpretable. The segments are useful hypotheses for marketing, not definitive customer identities, and require business validation."
    )
    doc.add_heading("References", level=1)
    doc.add_paragraph("1. Mall Customers dataset: https://github.com/kennedykwangari/Mall-Customer-Segmentation-Data")
    doc.add_paragraph("2. scikit-learn documentation, AgglomerativeClustering: https://scikit-learn.org/stable/modules/generated/sklearn.cluster.AgglomerativeClustering.html")
    doc.add_paragraph("3. scikit-learn documentation, silhouette_score: https://scikit-learn.org/stable/modules/generated/sklearn.metrics.silhouette_score.html")
    report_path = OUT / "Group12_Mini_Project_Report.docx"
    doc.save(report_path)

    summary = [
        "GROUP 12 - HIERARCHICAL AGGLOMERATIVE CLUSTERING",
        f"Dataset rows after cleaning: {len(clean)}",
        f"Original missing cells: {missing_before}",
        f"Original duplicate rows: {duplicates_before}",
        f"Selected clusters: {best_k}",
        f"Silhouette score: {silhouette_best:.4f}",
        f"Mean age: {age_mean:.2f}",
        f"Mean annual income (k$): {income_mean:.2f}",
        f"Mean spending score: {spending_mean:.2f}",
        f"Pearson correlation (income, spending): {income_spend_corr:.4f}",
        "",
        "Cluster profiles:",
        profile.to_string()
    ]
    (OUT / "results_summary.txt").write_text("\n".join(summary), encoding="utf-8")
    print("\n".join(summary))
    print(f"\nReport created: {report_path}")
    print(f"All outputs saved in: {OUT}")

if __name__ == "__main__":
    main()
